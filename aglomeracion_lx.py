
import io
import html

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# FUNCIONES AUXILIARES
# ============================================================

def numero(valor):
    """Presenta un valor con dos decimales."""
    return f"{valor:,.2f}"


def crear_excel(tabla, nombre_hoja="Balance de Masas"):
    """
    Crea un archivo Excel con datos numéricos reales
    y formato de dos decimales.
    """

    libro = Workbook()
    hoja = libro.active
    hoja.title = nombre_hoja[:31]

    encabezados = list(tabla.columns)

    hoja.append(encabezados)

    for fila in tabla.itertuples(index=False, name=None):
        hoja.append(list(fila))

    # Encabezado
    for celda in hoja[1]:
        celda.font = Font(
            bold=True,
            color="FFFFFF"
        )
        celda.fill = PatternFill(
            fill_type="solid",
            fgColor="17365D"
        )
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    hoja.row_dimensions[1].height = 25

    # Datos
    for fila in hoja.iter_rows(
        min_row=2,
        max_row=hoja.max_row
    ):
        for celda in fila:

            if isinstance(
                celda.value,
                (int, float)
            ):
                celda.number_format = '#,##0.00'

            celda.alignment = Alignment(
                vertical="center"
            )

    # Ancho automático aproximado
    for indice, columna in enumerate(
        hoja.columns,
        start=1
    ):
        longitud = max(
            len(str(celda.value or ""))
            for celda in columna
        )

        hoja.column_dimensions[
            get_column_letter(indice)
        ].width = min(
            max(longitud + 4, 18),
            48
        )

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    salida = io.BytesIO()

    libro.save(salida)
    salida.seek(0)

    return salida.getvalue()


def mostrar_tabla(tabla, clave, archivo):
    """
    Muestra una tabla con dos decimales
    y agrega su descarga en Excel.
    """

    columnas_numericas = [
        columna
        for columna in tabla.columns
        if pd.api.types.is_numeric_dtype(
            tabla[columna]
        )
    ]

    configuracion = {
        columna: st.column_config.NumberColumn(
            columna,
            format="%.2f"
        )
        for columna in columnas_numericas
    }

    st.dataframe(
        tabla,
        hide_index=True,
        use_container_width=True,
        column_config=configuracion
    )

    st.download_button(
        label="📥 Descargar tabla en Excel",
        data=crear_excel(tabla),
        file_name=archivo,
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        key=clave
    )


def tarjeta_fase(titulo, descripcion, color):

    titulo = html.escape(titulo)
    descripcion = html.escape(descripcion)

    contenido = f"""
    <div style="
        background:#1e2532;
        border:1px solid #334155;
        border-left:4px solid {color};
        border-radius:14px;
        padding:22px;
        min-height:180px;
        box-sizing:border-box;
    ">
        <h3 style="
            color:white;
            margin:0 0 14px 0;
        ">{titulo}</h3>

        <p style="
            color:#cbd5e1;
            line-height:1.6;
            margin:0;
        ">{descripcion}</p>
    </div>
    """

    st.markdown(
        contenido.replace("\n", ""),
        unsafe_allow_html=True
    )


# ============================================================
# DIAGRAMA DEL PROCESO DE AGLOMERACIÓN
# ============================================================

def mostrar_diagrama_aglomeracion(
    mineral_seco_tph,
    agua_mineral_tph,
    mineral_humedo_tph,
    agua_adicional_tph,
    acido_comercial_tph,
    mineral_aglomerado_tph,
    humedad_salida_pct
):

    valores = {
        "SECO": numero(mineral_seco_tph),
        "AGUA_MINERAL": numero(agua_mineral_tph),
        "HUMEDO": numero(mineral_humedo_tph),
        "AGUA_ADICIONAL": numero(agua_adicional_tph),
        "ACIDO": numero(acido_comercial_tph),
        "PRODUCTO": numero(mineral_aglomerado_tph),
        "HUMEDAD": numero(humedad_salida_pct)
    }

    svg = """
    <svg xmlns="http://www.w3.org/2000/svg"
         viewBox="0 0 1200 560"
         width="100%">

      <defs>
        <marker id="flecha"
                viewBox="0 0 10 10"
                refX="9"
                refY="5"
                markerWidth="7"
                markerHeight="7"
                orient="auto">
          <path d="M0 0 L10 5 L0 10 Z"
                fill="#94a3b8"/>
        </marker>

        <marker id="flechaVerde"
                viewBox="0 0 10 10"
                refX="9"
                refY="5"
                markerWidth="7"
                markerHeight="7"
                orient="auto">
          <path d="M0 0 L10 5 L0 10 Z"
                fill="#22c55e"/>
        </marker>

        <linearGradient id="tambor"
                        x1="0" y1="0"
                        x2="1" y2="1">
          <stop offset="0%"
                stop-color="#3b82f6"/>
          <stop offset="100%"
                stop-color="#1d4ed8"/>
        </linearGradient>
      </defs>

      <style>
        .caja {
          fill:#1e2532;
          stroke:#64748b;
          stroke-width:2;
        }

        .titulo {
          fill:#e2e8f0;
          font-family:Arial,sans-serif;
          font-size:18px;
          font-weight:bold;
          text-anchor:middle;
        }

        .valor {
          fill:white;
          font-family:Arial,sans-serif;
          font-size:23px;
          font-weight:bold;
          text-anchor:middle;
        }

        .linea {
          fill:none;
          stroke:#94a3b8;
          stroke-width:3;
          stroke-linecap:round;
          stroke-linejoin:round;
        }

        .flujo {
          fill:none;
          stroke:#94a3b8;
          stroke-width:3;
          stroke-linecap:round;
          stroke-linejoin:round;
          marker-end:url(#flecha);
        }
      </style>

      <!-- ENTRADAS SUPERIORES -->

      <rect x="400" y="25"
            width="230" height="90"
            rx="12"
            fill="#1e2532"
            stroke="#38bdf8"
            stroke-width="2"/>

      <text x="515" y="60" class="titulo">
        Agua adicional
      </text>

      <text x="515" y="94" class="valor">
        {{AGUA_ADICIONAL}} t/h
      </text>

      <rect x="715" y="25"
            width="230" height="90"
            rx="12"
            fill="#1e2532"
            stroke="#f59e0b"
            stroke-width="2"/>

      <text x="830" y="60" class="titulo">
        Ácido comercial
      </text>

      <text x="830" y="94" class="valor">
        {{ACIDO}} t/h
      </text>

      <!-- FLECHAS HACIA AGLOMERADOR -->

      <path d="M515 115 V195 H545 V250"
            class="flujo"/>

      <path d="M830 115 V180 H670 V250"
            class="flujo"/>

      <!-- MINERAL HÚMEDO -->

      <rect x="35" y="265"
            width="255" height="95"
            rx="12"
            class="caja"/>

      <text x="162.5" y="301" class="titulo">
        Mineral húmedo
      </text>

      <text x="162.5" y="337" class="valor">
        {{HUMEDO}} t/h
      </text>

      <path d="M290 312 H405"
            class="flujo"/>

      <!-- AGLOMERADOR -->

      <rect x="415" y="255"
            width="310" height="115"
            rx="28"
            fill="url(#tambor)"
            stroke="#93c5fd"
            stroke-width="3"/>

      <ellipse cx="725" cy="312.5"
               rx="25" ry="57.5"
               fill="#1d4ed8"
               stroke="#93c5fd"
               stroke-width="2"/>

      <text x="560" y="323"
            text-anchor="middle"
            fill="white"
            font-family="Arial,sans-serif"
            font-size="27"
            font-weight="bold">
        Aglomerador
      </text>

      <!-- SALIDA -->

      <path d="M750 312 H865"
            fill="none"
            stroke="#22c55e"
            stroke-width="3"
            marker-end="url(#flechaVerde)"/>

      <rect x="875" y="265"
            width="290" height="95"
            rx="12"
            fill="#1e2532"
            stroke="#22c55e"
            stroke-width="2"/>

      <text x="1020" y="301" class="titulo">
        Mineral aglomerado
      </text>

      <text x="1020" y="337" class="valor">
        {{PRODUCTO}} t/h
      </text>

      <text x="1020" y="405"
            fill="#cbd5e1"
            text-anchor="middle"
            font-family="Arial,sans-serif"
            font-size="17">
        Humedad de salida
      </text>

      <text x="1020" y="445"
            fill="#22c55e"
            text-anchor="middle"
            font-family="Arial,sans-serif"
            font-size="32"
            font-weight="bold">
        {{HUMEDAD}}%
      </text>

      <!-- MINERAL SECO -->

      <rect x="25" y="470"
            width="185" height="75"
            rx="12"
            class="caja"/>

      <text x="117.5" y="499" class="titulo">
        Mineral seco
      </text>

      <text x="117.5" y="528" class="valor">
        {{SECO}} t/h
      </text>

      <!-- SIGNO MÁS -->

      <text x="245" y="520"
            fill="white"
            text-anchor="middle"
            font-family="Arial,sans-serif"
            font-size="42"
            font-weight="bold">
        +
      </text>

      <!-- AGUA DEL MINERAL -->

      <rect x="280" y="470"
            width="195" height="75"
            rx="12"
            fill="#1e2532"
            stroke="#38bdf8"
            stroke-width="2"/>

      <text x="377.5" y="499" class="titulo">
        Agua del mineral
      </text>

      <text x="377.5" y="528" class="valor">
        {{AGUA_MINERAL}} t/h
      </text>

      <!-- UNIÓN DE AMBOS COMPONENTES -->

      <path d="M117.5 470 V425 H162.5"
            class="linea"/>

      <path d="M377.5 470 V425 H162.5"
            class="linea"/>

      <path d="M162.5 425 V365"
            class="flujo"/>

    </svg>
    """

    for clave, valor in valores.items():
        svg = svg.replace(
            "{{" + clave + "}}",
            html.escape(valor)
        )

    pagina = f"""
    <!DOCTYPE html>
    <html lang="es">
    <head>
      <meta charset="UTF-8">
      <style>
        html, body {{
          margin:0;
          padding:0;
          background:transparent;
          overflow:hidden;
        }}

        svg {{
          display:block;
          width:100%;
          height:auto;
        }}
      </style>
    </head>
    <body>
      {svg}
    </body>
    </html>
    """

    components.html(
        pagina,
        height=570,
        scrolling=False
    )


# ============================================================
# 1. PROCESOS HIDROMETALÚRGICOS
# ============================================================

def mostrar_procesos_hidro():

    st.title("💧 Procesos Hidrometalúrgicos")

    st.write(
        "Conoce las principales etapas utilizadas "
        "para la recuperación de cobre a partir "
        "de minerales oxidados."
    )

    st.markdown("---")
    st.markdown("### Etapas del Proceso")

    fases = [
        (
            "1. Aglomeración",
            "Mezcla del mineral con agua y ácido "
            "sulfúrico cuando corresponde. "
            "Favorece la permeabilidad del lecho.",
            "#a855f7"
        ),
        (
            "2. Lixiviación",
            "Disolución del cobre mediante "
            "soluciones acuosas ácidas, "
            "obteniéndose una solución rica "
            "en cobre denominada PLS.",
            "#22c55e"
        ),
        (
            "3. Extracción por Solventes",
            "Transferencia selectiva del cobre "
            "hacia una fase orgánica y "
            "posterior reextracción.",
            "#3b82f6"
        ),
        (
            "4. Electroobtención",
            "Recuperación de cobre metálico "
            "mediante corriente eléctrica.",
            "#f59e0b"
        )
    ]

    for i in range(0, len(fases), 2):

        columnas = st.columns(2)

        for columna, fase in zip(
            columnas,
            fases[i:i + 2]
        ):
            with columna:
                tarjeta_fase(*fase)

        st.write("")


# ============================================================
# 2. CALCULADORA DE AGLOMERACIÓN
# ============================================================

def mostrar_aglomeracion():

    st.title(
        "🧮 Simulador y Diagrama de Balance de Masas: "
        "Aglomeración"
    )

    st.write(
        "Calcula las corrientes de entrada y salida "
        "del aglomerador a partir del flujo de mineral, "
        "la humedad inicial y las adiciones de agua "
        "y ácido sulfúrico."
    )

    st.markdown("---")
    st.markdown("### 📥 Parámetros de Operación")

    col1, col2, col3, col4 = st.columns(4)

    with col1:

        tipo_mineral = st.radio(
            "Base de mineral ingresado",
            [
                "Mineral Húmedo (ton/h)",
                "Mineral Seco (ton/h)"
            ],
            key="agl_tipo_mineral"
        )

        val_mineral = st.number_input(
            "Valor de mineral ingresado",
            min_value=0.0,
            value=736.5,
            step=10.0,
            format="%.2f",
            key="agl_val_mineral"
        )

    with col2:

        humedad_inicial = st.number_input(
            "Humedad inicial del mineral [%]",
            min_value=0.0,
            max_value=99.9,
            value=2.0,
            step=0.1,
            format="%.2f",
            key="agl_humedad_inicial"
        )

    with col3:

        unidad_agua = st.radio(
            "Unidad de agua adicional",
            [
                "kg/h",
                "ton/h",
                "ton/min"
            ],
            key="agl_unidad_agua"
        )

        val_agua = st.number_input(
            "Agua adicional",
            min_value=0.0,
            value=0.5,
            step=0.1,
            format="%.2f",
            key="agl_val_agua"
        )

    with col4:

        unidad_acido = st.radio(
            "Unidad de ácido H₂SO₄",
            [
                "kg/h",
                "ton/h",
                "kg/ton de seco"
            ],
            key="agl_unidad_acido"
        )

        val_acido = st.number_input(
            "Ácido comercial",
            min_value=0.0,
            value=33500.0,
            step=500.0,
            format="%.2f",
            key="agl_val_acido"
        )

    pureza_acido = st.number_input(
        "Concentración del ácido comercial [% H₂SO₄]",
        min_value=0.1,
        max_value=100.0,
        value=98.0,
        step=0.5,
        format="%.2f",
        key="agl_pureza_acido"
    )

    # ========================================================
    # BALANCE DE MASAS
    # ========================================================

    # Agua adicional a t/h
    if unidad_agua == "kg/h":
        agua_adicional_tph = val_agua / 1000.0

    elif unidad_agua == "ton/min":
        agua_adicional_tph = val_agua * 60.0

    else:
        agua_adicional_tph = val_agua

    # Mineral seco, húmedo y agua del mineral
    fraccion_humedad = humedad_inicial / 100.0

    if tipo_mineral == "Mineral Húmedo (ton/h)":

        mineral_humedo_tph = val_mineral

        agua_mineral_tph = (
            mineral_humedo_tph * fraccion_humedad
        )

        mineral_seco_tph = (
            mineral_humedo_tph - agua_mineral_tph
        )

    else:

        mineral_seco_tph = val_mineral

        mineral_humedo_tph = (
            mineral_seco_tph
            / (1.0 - fraccion_humedad)
        )

        agua_mineral_tph = (
            mineral_humedo_tph - mineral_seco_tph
        )

    # Ácido comercial a t/h
    if unidad_acido == "kg/h":

        acido_comercial_tph = (
            val_acido / 1000.0
        )

    elif unidad_acido == "kg/ton de seco":

        acido_comercial_tph = (
            val_acido
            * mineral_seco_tph
            / 1000.0
        )

    else:

        acido_comercial_tph = val_acido

    # Composición del ácido comercial
    fraccion_acido = pureza_acido / 100.0

    acido_puro_tph = (
        acido_comercial_tph
        * fraccion_acido
    )

    agua_acido_tph = (
        acido_comercial_tph
        * (1.0 - fraccion_acido)
    )

    # Producto
    mineral_aglomerado_tph = (
        mineral_humedo_tph
        + agua_adicional_tph
        + acido_comercial_tph
    )

    agua_total_tph = (
        agua_mineral_tph
        + agua_adicional_tph
        + agua_acido_tph
    )

    masa_no_agua_tph = (
        mineral_seco_tph
        + acido_puro_tph
    )

    # Humedad de salida
    if mineral_aglomerado_tph > 0:

        humedad_salida_pct = (
            agua_total_tph
            / mineral_aglomerado_tph
            * 100.0
        )

    else:

        humedad_salida_pct = 0.0

    # Verificación
    entrada_total_tph = (
        mineral_humedo_tph
        + agua_adicional_tph
        + acido_comercial_tph
    )

    salida_total_tph = (
        masa_no_agua_tph
        + agua_total_tph
    )

    error_balance = abs(
        entrada_total_tph
        - salida_total_tph
    )

    # ========================================================
    # DIAGRAMA
    # ========================================================

    st.markdown("---")
    st.markdown(
        "### 🔄 Diagrama del Proceso de Aglomeración"
    )

    mostrar_diagrama_aglomeracion(
        mineral_seco_tph=mineral_seco_tph,
        agua_mineral_tph=agua_mineral_tph,
        mineral_humedo_tph=mineral_humedo_tph,
        agua_adicional_tph=agua_adicional_tph,
        acido_comercial_tph=acido_comercial_tph,
        mineral_aglomerado_tph=mineral_aglomerado_tph,
        humedad_salida_pct=humedad_salida_pct
    )

    st.caption(
        "El mineral seco y el agua del mineral "
        "forman el mineral húmedo. "
        "El agua adicional y el ácido comercial "
        "ingresan al aglomerador por separado."
    )

    # ========================================================
    # RESUMEN DE MASAS
    # ========================================================

    st.markdown("---")
    st.markdown(
        "### 📋 Resumen de Masas del Proceso"
    )

    r1, r2, r3, r4 = st.columns(4)

    r1.metric(
        "Mineral seco",
        f"{numero(mineral_seco_tph)} t/h"
    )

    r2.metric(
        "Agua del mineral",
        f"{numero(agua_mineral_tph)} t/h"
    )

    r3.metric(
        "Mineral húmedo",
        f"{numero(mineral_humedo_tph)} t/h"
    )

    r4.metric(
        "Mineral aglomerado",
        f"{numero(mineral_aglomerado_tph)} t/h"
    )

    # ========================================================
    # BALANCE DETALLADO DE MASAS
    # ========================================================

    st.markdown("---")
    st.markdown(
        "### 📊 Balance Detallado de Masas"
    )

    tabla_balance = pd.DataFrame({
        "Corriente": [
            "Mineral seco",
            "Agua del mineral",
            "Agua adicional",
            "Ácido sulfúrico puro",
            "Agua contenida en el ácido comercial",
            "Mineral aglomerado"
        ],
        "Flujo [t/h]": [
            mineral_seco_tph,
            agua_mineral_tph,
            agua_adicional_tph,
            acido_puro_tph,
            agua_acido_tph,
            mineral_aglomerado_tph
        ]
    })

    mostrar_tabla(
        tabla=tabla_balance,
        clave="descarga_balance_aglomeracion",
        archivo="balance_masas_aglomeracion.xlsx"
    )

    # ========================================================
    # VERIFICACIÓN DEL BALANCE
    # ========================================================

    st.markdown(
        "### ✅ Verificación del Balance"
    )

    if error_balance < 1e-7:

        st.success(
            "El balance de masas está cerrado."
        )

    else:

        st.error(
            "Existe una diferencia de "
            f"{numero(error_balance)} t/h."
        )

    st.info(
        "Humedad estimada del producto: "
        f"{numero(humedad_salida_pct)}%."
    )

    st.caption(
        "Modelo simplificado sin pérdidas, "
        "evaporación ni reacciones que alteren "
        "las masas de los componentes. "
        "La humedad se expresa como porcentaje "
        "másico de agua."
    )


# ============================================================
# 3. PREPARACIÓN DE SOLUCIÓN LX
# ============================================================

def mostrar_preparacion_lx():

    st.title(
        "🧪 Preparación de Solución para Lixiviación"
    )

    st.write(
        "Estimación de la dosificación de ácido "
        "sulfúrico comercial y agua para preparar "
        "una solución de lixiviación a partir "
        "de un pH objetivo."
    )

    st.warning(
        "Este cálculo es una aproximación ideal. "
        "En soluciones reales, el pH no permite "
        "determinar por sí solo la concentración "
        "total de ácido sulfúrico."
    )

    st.markdown("---")
    st.markdown(
        "### 📥 Parámetros de Operación"
    )

    col1, col2 = st.columns(2)

    with col1:

        q_solucion = st.number_input(
            "Caudal de solución [m³/h]",
            min_value=0.1,
            value=1000.0,
            step=10.0,
            format="%.2f",
            key="lx_q_solucion"
        )

        ph_objetivo = st.number_input(
            "pH objetivo",
            min_value=0.0,
            max_value=7.0,
            value=0.50,
            step=0.1,
            format="%.2f",
            key="lx_ph_objetivo"
        )

    with col2:

        pureza_acido = st.number_input(
            "Pureza del ácido comercial [%]",
            min_value=50.0,
            max_value=100.0,
            value=98.0,
            step=0.5,
            format="%.2f",
            key="lx_pureza_acido"
        )

        densidad_acido = st.number_input(
            "Densidad del ácido comercial [t/m³]",
            min_value=1.0,
            max_value=2.5,
            value=1.84,
            step=0.01,
            format="%.2f",
            key="lx_densidad_acido"
        )

    # ========================================================
    # CÁLCULOS
    # ========================================================

    concentracion_h = (
        10.0 ** (-ph_objetivo)
    )

    masa_molar_acido = 98.079

    # Aproximación ideal de dos protones
    # por mol de ácido sulfúrico.
    acido_puro_g_l = (
        concentracion_h
        * masa_molar_acido
        / 2.0
    )

    acido_puro_kg_h = (
        acido_puro_g_l
        * q_solucion
    )

    fraccion_pureza = (
        pureza_acido / 100.0
    )

    acido_comercial_kg_h = (
        acido_puro_kg_h
        / fraccion_pureza
    )

    acido_comercial_t_h = (
        acido_comercial_kg_h
        / 1000.0
    )

    volumen_acido_m3_h = (
        acido_comercial_t_h
        / densidad_acido
    )

    # Supuesto de volúmenes aditivos
    agua_m3_h = (
        q_solucion
        - volumen_acido_m3_h
    )

    if agua_m3_h < 0:

        st.error(
            "El volumen calculado de ácido "
            "supera el caudal total de solución."
        )

        return

    # ========================================================
    # RESULTADOS
    # ========================================================

    st.markdown("---")
    st.markdown(
        "### 📊 Resultados de Dosificación"
    )

    r1, r2 = st.columns(2)

    r1.metric(
        "Ácido comercial",
        f"{numero(volumen_acido_m3_h)} m³/h"
    )

    r2.metric(
        "Agua estimada",
        f"{numero(agua_m3_h)} m³/h"
    )

    r3, r4 = st.columns(2)

    r3.metric(
        "Ácido puro",
        f"{numero(acido_puro_kg_h)} kg/h"
    )

    r4.metric(
        "Concentración estimada",
        f"{numero(acido_puro_g_l)} g/L"
    )

    # ========================================================
    # TABLA DE RESULTADOS LX
    # ========================================================

    st.markdown("---")
    st.markdown(
        "### 📋 Resumen de Preparación de Solución LX"
    )

    tabla_lx = pd.DataFrame({
        "Parámetro": [
            "Caudal de solución",
            "pH objetivo",
            "Pureza del ácido comercial",
            "Densidad del ácido comercial",
            "Concentración estimada de ácido puro",
            "Masa de ácido puro",
            "Masa de ácido comercial",
            "Volumen de ácido comercial",
            "Caudal de agua estimado"
        ],
        "Valor": [
            q_solucion,
            ph_objetivo,
            pureza_acido,
            densidad_acido,
            acido_puro_g_l,
            acido_puro_kg_h,
            acido_comercial_kg_h,
            volumen_acido_m3_h,
            agua_m3_h
        ],
        "Unidad": [
            "m³/h",
            "pH",
            "%",
            "t/m³",
            "g/L",
            "kg/h",
            "kg/h",
            "m³/h",
            "m³/h"
        ]
    })

    mostrar_tabla(
        tabla=tabla_lx,
        clave="descarga_preparacion_lx",
        archivo="preparacion_solucion_lx.xlsx"
    )

    st.info(
        "Para preparar aproximadamente "
        f"{numero(q_solucion)} m³/h de solución, "
        "el modelo estima "
        f"{numero(volumen_acido_m3_h)} m³/h "
        "de ácido comercial y "
        f"{numero(agua_m3_h)} m³/h de agua."
    )

    st.caption(
        "Supuestos: disociación completa de "
        "los dos protones del H₂SO₄, "
        "actividad igual a concentración "
        "y volúmenes aditivos."
    )

    st.warning(
        "Para soluciones reales se recomienda "
        "trabajar con concentración de H₂SO₄ "
        "en g/L y mediciones de acidez libre. "
        "Por seguridad, siempre se agrega "
        "el ácido al agua, nunca al revés."
    )
